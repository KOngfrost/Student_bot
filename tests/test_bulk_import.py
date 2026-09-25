"""
Тесты массового импорта и экспорта данных (Excel/CSV) для Частых вопросов и Базы знаний,
а также работы обновлённых хендлеров бота с инструкциями.
"""

import asyncio
import contextlib
import io
import threading
import time
from unittest.mock import AsyncMock

import openpyxl
import pytest

import core.bulk_import as bulk_import_module
from bots.vk.handlers.faq import (
    faq_handler,
)
from bots.vk.handlers.knowledge import (
    knowledge_base_handler,
)
from core.bulk_import import (
    MAX_IMPORT_FILE_SIZE,
    MAX_IMPORT_ROWS,
    ImportLimitError,
    export_faq_xlsx,
    export_faq_xlsx_async,
    export_knowledge_xlsx,
    export_knowledge_xlsx_async,
    generate_faq_template_csv,
    generate_faq_template_csv_async,
    generate_faq_template_xlsx,
    generate_faq_template_xlsx_async,
    generate_knowledge_template_csv,
    generate_knowledge_template_csv_async,
    generate_knowledge_template_xlsx,
    generate_knowledge_template_xlsx_async,
    parse_faq_rows,
    parse_faq_rows_async,
    parse_file_or_text,
    parse_file_or_text_async,
    parse_knowledge_rows,
    parse_knowledge_rows_async,
)


def test_generate_faq_template_xlsx():
    data = generate_faq_template_xlsx()
    assert isinstance(data, bytes)
    assert len(data) > 100
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.title == "Частые вопросы"
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Вопрос"
    assert rows[0][1] == "Ответ"
    assert len(rows) > 1


def test_generate_faq_template_csv():
    data = generate_faq_template_csv()
    assert "Вопрос;Ответ" in data
    assert "Сколько всего общежитий" in data


def test_generate_knowledge_template_xlsx():
    data = generate_knowledge_template_xlsx()
    assert isinstance(data, bytes)
    assert len(data) > 100
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.title == "База знаний"
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Ключевые слова"


def test_generate_knowledge_template_csv():
    data = generate_knowledge_template_csv()
    assert "Ключевые слова;Материал / Ответ" in data
    assert "сообщество" in data


def test_parse_faq_rows():
    raw = [
        ["Вопрос", "Ответ", "Отдел"],
        ["Где получить пропуск?", "На вахте общежития", "Жилищно-бытовой"],
        ["", "", ""],
        ["Как записаться?", "В боте", "Культмасс"],
    ]
    parsed = parse_faq_rows(raw)
    assert len(parsed) == 2
    assert parsed[0]["question"] == "Где получить пропуск?"
    assert parsed[0]["answer"] == "На вахте общежития"
    assert parsed[0]["department"] == "Жилищно-бытовой"
    assert parsed[1]["question"] == "Как записаться?"


def test_parse_knowledge_rows():
    raw = [
        ["Ключевые слова", "Ответ / Материал", "Отдел"],
        ["сантехника, кран", "Звоните мастеру", "Жилбыт"],
        ["мероприятия", "Смотрите афишу ВК", "Культмасс"],
    ]
    parsed = parse_knowledge_rows(raw)
    assert len(parsed) == 2
    assert parsed[0]["keywords"] == "сантехника, кран"
    assert parsed[0]["answer"] == "Звоните мастеру"


def test_parse_file_or_text_csv():
    csv_text = "Вопрос;Ответ;Отдел\nКак заселиться?;С паспортом;Жилбыт\n"
    raw = parse_file_or_text(None, None, csv_text)
    assert len(raw) == 2
    parsed = parse_faq_rows(raw)
    assert len(parsed) == 1
    assert parsed[0]["question"] == "Как заселиться?"


# ---------------------------------------------------------------------------
# SEC-10: защита от Zip/XML-бомб — лимиты размера файла и числа строк.
# ---------------------------------------------------------------------------


def test_parse_file_rejects_oversized_file():
    """Файл больше MAX_IMPORT_FILE_SIZE отклоняется до разбора (SEC-10)."""
    with pytest.raises(ImportLimitError) as exc_info:
        parse_file_or_text(b"x" * (MAX_IMPORT_FILE_SIZE + 1), "bomb.xlsx")
    assert "5 МБ" in str(exc_info.value)


def test_parse_text_rejects_too_many_rows():
    """Текстовая вставка с числом строк больше лимита прерывается (SEC-10)."""
    rows = "\n".join(f"Вопрос {i}?;Ответ {i}" for i in range(MAX_IMPORT_ROWS + 5))
    with pytest.raises(ImportLimitError) as exc_info:
        parse_file_or_text(None, None, rows)
    assert "строк" in str(exc_info.value)


def test_parse_xlsx_rejects_too_many_rows():
    """xlsx с числом строк больше лимита прерывается до исчерпания памяти (SEC-10)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Вопрос", "Ответ"])
    for i in range(MAX_IMPORT_ROWS + 5):
        ws.append([f"Вопрос {i}?", f"Ответ {i}"])
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(ImportLimitError):
        parse_file_or_text(buf.getvalue(), "bomb.xlsx")


def test_parse_within_limits_still_works():
    """Вход в пределах лимитов разбирается как раньше (регрессия SEC-10)."""
    rows = "\n".join(f"Вопрос {i}?;Ответ {i}" for i in range(50))
    raw = parse_file_or_text(None, None, rows)
    assert len(raw) == 50


def test_export_faq_and_knowledge_xlsx():
    class DummyDept:
        name = "Информационный"

    class DummyFAQ:
        id = 1
        department = DummyDept()
        question = "Тестовый вопрос?"
        final_answer = "Тестовый ответ"

    faq_bytes = export_faq_xlsx([DummyFAQ()])
    assert len(faq_bytes) > 100

    class DummyKB:
        id = 1
        department = DummyDept()
        keywords = "тест, справка"
        answer = "Информация о тесте"

    kb_bytes = export_knowledge_xlsx([DummyKB()])
    assert len(kb_bytes) > 100


# ---------------------------------------------------------------------------
# Асинхронный API: обёртки *_async обязаны выполнять синхронные ядра openpyxl
# в пуле потоков (asyncio.to_thread) и не замораживать event loop (Ошибка #11).
# ---------------------------------------------------------------------------


def _xlsx_rows(data: bytes) -> list[tuple]:
    """Значения ячеек первого листа xlsx.

    Сравнивать файлы побайтово нельзя: openpyxl пишет в docProps метку времени
    создания книги, поэтому две одинаковые выгрузки отличаются в байтах.
    """
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return list(wb.active.iter_rows(values_only=True))


def _dummy_faq():
    class _Dept:
        name = "Информационный"

    class _FAQ:
        id = 1
        department = _Dept()
        question = "Тестовый вопрос?"
        final_answer = "Тестовый ответ"

    return _FAQ()


def _dummy_knowledge():
    class _Dept:
        name = "Информационный"

    class _KB:
        id = 1
        department = _Dept()
        keywords = "тест, справка"
        answer = "Информация о тесте"

    return _KB()


@pytest.mark.asyncio
async def test_async_wrappers_return_same_as_sync_cores():
    """Каждая *_async обёртка возвращает ровно то же, что синхронное ядро."""
    raw = [
        ["Вопрос", "Ответ", "Отдел"],
        ["Где коворкинг?", "На 3 этаже", "Информационный"],
        ["", "", ""],
    ]
    csv_text = "Вопрос;Ответ;Отдел\nКак заселиться?;С паспортом;Жилбыт\n"

    assert _xlsx_rows(await generate_faq_template_xlsx_async()) == _xlsx_rows(
        generate_faq_template_xlsx()
    )
    assert await generate_faq_template_csv_async() == generate_faq_template_csv()
    assert _xlsx_rows(await generate_knowledge_template_xlsx_async()) == _xlsx_rows(
        generate_knowledge_template_xlsx()
    )
    assert await generate_knowledge_template_csv_async() == generate_knowledge_template_csv()

    assert await parse_faq_rows_async(raw) == parse_faq_rows(raw)
    assert await parse_knowledge_rows_async(raw) == parse_knowledge_rows(raw)
    assert await parse_file_or_text_async(None, None, csv_text) == parse_file_or_text(
        None, None, csv_text
    )

    assert _xlsx_rows(await export_faq_xlsx_async([_dummy_faq()])) == _xlsx_rows(
        export_faq_xlsx([_dummy_faq()])
    )
    assert _xlsx_rows(await export_knowledge_xlsx_async([_dummy_knowledge()])) == _xlsx_rows(
        export_knowledge_xlsx([_dummy_knowledge()])
    )


@pytest.mark.asyncio
async def test_async_wrappers_run_sync_core_in_worker_thread(monkeypatch):
    """Синхронное ядро исполняется в рабочем потоке, а не в потоке event loop."""
    loop_thread = threading.get_ident()
    core_threads: set[int] = set()

    def _fake_sync_core(raw_rows):
        core_threads.add(threading.get_ident())
        return [{"question": "q", "answer": "a", "department": ""}]

    monkeypatch.setattr(bulk_import_module, "parse_faq_rows", _fake_sync_core)

    result = await parse_faq_rows_async([["Вопрос", "Ответ"]])

    assert result == [{"question": "q", "answer": "a", "department": ""}]
    assert len(core_threads) == 1
    assert core_threads.pop() != loop_thread


@pytest.mark.asyncio
async def test_heavy_import_does_not_block_event_loop(monkeypatch):
    """Тяжёлый разбор в потоке не замораживает event loop и идёт параллельно."""
    sync_core_seconds = 0.2
    parsed_rows = [["Вопрос", "Ответ"], ["Как заселиться?", "С паспортом"]]

    def _slow_sync_core(file_bytes, filename, text_content=None):
        time.sleep(sync_core_seconds)
        return parsed_rows

    monkeypatch.setattr(bulk_import_module, "parse_file_or_text", _slow_sync_core)

    ticks = 0

    async def _ticker():
        nonlocal ticks
        while True:
            await asyncio.sleep(0.01)
            ticks += 1

    ticker = asyncio.create_task(_ticker())
    loop = asyncio.get_running_loop()
    began = loop.time()
    try:
        results = await asyncio.gather(
            parse_file_or_text_async(None, None, "a"),
            parse_file_or_text_async(None, None, "b"),
            parse_file_or_text_async(None, None, "c"),
        )
        elapsed = loop.time() - began
    finally:
        ticker.cancel()
        with contextlib.suppress(asyncio.CancelledError):
            await ticker

    assert results == [parsed_rows, parsed_rows, parsed_rows]
    # Если бы event loop был занят синхронным ядром, тикер не проснулся бы ни разу.
    assert ticks >= 1
    # Три ядра по 0.2 с идут в пуле потоков параллельно, а не последовательно.
    assert elapsed < sync_core_seconds * 2


@pytest.mark.asyncio
async def test_faq_handler_returns_instructions(db_session_maker):
    msg = AsyncMock()
    msg.text = "Частые вопросы"
    msg.from_id = 12345

    await faq_handler(msg)

    assert msg.answer.called
    sent_text = msg.answer.call_args[0][0]
    assert "Инструкция по разделу «Частые вопросы»" in sent_text
    assert "Как пользоваться" in sent_text


@pytest.mark.asyncio
async def test_knowledge_base_handler_returns_instructions(db_session_maker):
    msg = AsyncMock()
    msg.text = "База знаний"
    msg.from_id = 12345

    await knowledge_base_handler(msg)

    assert msg.answer.called
    sent_text = msg.answer.call_args[0][0]
    assert "Инструкция по Базе знаний" in sent_text
    assert "Как пользоваться" in sent_text


def _login(client):
    import re

    login_page = client.get("/auth/login")
    csrf_token = re.search(r'name="csrf_token" value="([^"]+)"', login_page.text).group(1)
    client.post(
        "/auth/login",
        data={"username": "testadmin", "password": "test_password_123", "csrf_token": csrf_token},
        follow_redirects=False,
    )


def test_faq_web_endpoints(web_client):
    _login(web_client)

    resp = web_client.get("/faq/template.xlsx")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]

    resp = web_client.get("/faq/template.csv")
    assert resp.status_code == 200
    assert "Вопрос;Ответ" in resp.text

    resp = web_client.get("/faq/export.xlsx")
    assert resp.status_code == 200


def test_knowledge_web_endpoints(web_client):
    _login(web_client)

    resp = web_client.get("/knowledge/template.xlsx")
    assert resp.status_code == 200

    resp = web_client.get("/knowledge/template.csv")
    assert resp.status_code == 200

    resp = web_client.get("/knowledge/export.xlsx")
    assert resp.status_code == 200


def test_faq_and_knowledge_bulk_import_post(web_client):
    import re

    _login(web_client)

    # FAQ import
    faq_page = web_client.get("/faq/")
    csrf = re.search(r'name="csrf_token" value="([^"]+)"', faq_page.text).group(1)
    csv_text = "Вопрос;Ответ\nГде коворкинг?;На 3 этаже главного здания\n"
    resp = web_client.post(
        "/faq/import",
        data={"text_data": csv_text, "csrf_token": csrf},
        follow_redirects=True,
    )
    assert resp.status_code == 200
    assert "Успешно загружено" in resp.text

    # Knowledge import
    kb_page = web_client.get("/knowledge/")
    csrf_kb = re.search(r'name="csrf_token" value="([^"]+)"', kb_page.text).group(1)
    kb_csv = "Ключевые слова;Ответ\nспортзал, тренировки;Спортзал открыт с 8 до 22\n"
    resp_kb = web_client.post(
        "/knowledge/import",
        data={"text_data": kb_csv, "csrf_token": csrf_kb},
        follow_redirects=True,
    )
    assert resp_kb.status_code == 200
    assert "Успешно загружено" in resp_kb.text
