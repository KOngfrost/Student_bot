"""
Модуль массового импорта и экспорта данных (Excel / CSV) для Частых вопросов и Базы знаний.

Двухслойный API:

1. Синхронные ядра (``generate_*``, ``parse_*``, ``export_*``) — ресурсоёмкие
   вызовы openpyxl (``load_workbook`` / ``Workbook.save``) и модуля ``csv``.
   Предназначены только для синхронных потребителей (тесты, CLI-скрипты) и
   сами по себе блокируют поток вызова.
2. Асинхронные обёртки (те же имена с суффиксом ``_async``) — публичный API для
   FastAPI-обработчиков и VK-бота: каждое синхронное ядро выполняется в пуле
   потоков через ``asyncio.to_thread``, поэтому разбор больших файлов не
   блокирует event loop (Ошибка #11).

Правило: из async-контекста вызывать только ``*_async`` функции.
"""

import asyncio
import csv
import io
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


def _decode_bytes(file_bytes: bytes) -> str:
    """Безопасное декодирование байт с подбором кодировок."""
    for enc in ("utf-8-sig", "utf-8", "cp1251", "latin1"):
        try:
            return file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _detect_delimiter(text: str) -> str:
    """Определение разделителя CSV."""
    first_lines = text.strip().split("\n")[:5]
    sample = "\n".join(first_lines)
    if ";" in sample and sample.count(";") >= sample.count(","):
        return ";"
    if "\t" in sample:
        return "\t"
    return ","


def generate_faq_template_xlsx() -> bytes:
    """Сгенерировать красивый Excel-шаблон для загрузки Частых вопросов."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Частые вопросы"

    headers = ["Вопрос", "Ответ", "Отдел (необязательно)"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    samples = [
        [
            "Сколько всего общежитий в политехе?",
            "В Студгородке СПбПУ 16 общежитий, расположенных в шаговой доступности от учебных корпусов.",
            "Жилищно-бытовой",
        ],
        [
            "Как записаться на мероприятие Культмасса?",
            "Запись на мероприятия доступна через раздел «Мероприятия» в боте или в группе ВК.",
            "Культурно-массовый",
        ],
        [
            "Где заказать справку об обучении?",
            "Справку об обучении можно заказать через личный кабинет студента или в Единой дирекции.",
            "Информационный",
        ],
    ]
    for row in samples:
        ws.append(row)

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 65
    ws.column_dimensions["C"].width = 25

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def generate_faq_template_csv() -> str:
    """Сгенерировать CSV-шаблон для Частых вопросов."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Вопрос", "Ответ", "Отдел (необязательно)"])
    writer.writerow(
        [
            "Сколько всего общежитий в политехе?",
            "В Студгородке СПбПУ 16 общежитий рядом с кампусом.",
            "Жилищно-бытовой",
        ]
    )
    writer.writerow(
        [
            "Как записаться на мероприятие?",
            "Запись на мероприятия доступна в разделе «Мероприятия» в боте.",
            "Культурно-массовый",
        ]
    )
    return output.getvalue()


def generate_knowledge_template_xlsx() -> bytes:
    """Сгенерировать Excel-шаблон для загрузки Базы знаний."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "База знаний"

    headers = ["Ключевые слова", "Материал / Ответ", "Отдел (необязательно)"]
    ws.append(headers)

    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    samples = [
        [
            "сообщество, группа вк, культмасс",
            "Мы публикуем все мероприятия в сообществе: https://vk.ru/kultmass_oss_spbpu",
            "Культурно-массовый",
        ],
        [
            "сантехника, кран, протечка, вызов мастера",
            "Для вызова мастера оставьте заявку в журнале на вахте или обратитесь в Студсовет.",
            "Жилищно-бытовой",
        ],
    ]
    for row in samples:
        ws.append(row)

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 65
    ws.column_dimensions["C"].width = 25

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def generate_knowledge_template_csv() -> str:
    """Сгенерировать CSV-шаблон для Базы знаний."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Ключевые слова", "Материал / Ответ", "Отдел (необязательно)"])
    writer.writerow(
        [
            "сообщество, группа вк, культмасс",
            "Мы публикуем мероприятия: https://vk.ru/kultmass_oss_spbpu",
            "Культурно-массовый",
        ]
    )
    writer.writerow(
        [
            "сантехника, кран, вызов",
            "Оставьте запись на вахте или подайте заявку через бот.",
            "Жилищно-бытовой",
        ]
    )
    return output.getvalue()


def parse_faq_rows(raw_rows: list[list[str]]) -> list[dict[str, str]]:
    """Извлечь строки вопросов и ответов из сырых списков колонок."""
    if not raw_rows:
        return []

    header_row = [str(c or "").strip().lower() for c in raw_rows[0]]
    q_idx, a_idx, d_idx = -1, -1, -1

    for idx, col in enumerate(header_row):
        if "вопрос" in col or "question" in col:
            q_idx = idx
        elif "ответ" in col or "answer" in col:
            a_idx = idx
        elif "отдел" in col or "dept" in col or "department" in col:
            d_idx = idx

    start_row = 1 if (q_idx != -1 or a_idx != -1) else 0
    if q_idx == -1:
        q_idx = 0
    if a_idx == -1:
        a_idx = 1 if len(raw_rows[0]) > 1 else -1

    results = []
    for row in raw_rows[start_row:]:
        if not row:
            continue
        q = str(row[q_idx] if q_idx < len(row) and row[q_idx] is not None else "").strip()
        a = str(
            row[a_idx] if a_idx != -1 and a_idx < len(row) and row[a_idx] is not None else ""
        ).strip()
        d = str(
            row[d_idx] if d_idx != -1 and d_idx < len(row) and row[d_idx] is not None else ""
        ).strip()

        if q and a:
            results.append({"question": q, "answer": a, "department": d})

    return results


def parse_knowledge_rows(raw_rows: list[list[str]]) -> list[dict[str, str]]:
    """Извлечь записи базы знаний из сырых строк колонок."""
    if not raw_rows:
        return []

    header_row = [str(c or "").strip().lower() for c in raw_rows[0]]
    k_idx, a_idx, d_idx = -1, -1, -1

    for idx, col in enumerate(header_row):
        if "ключ" in col or "keyword" in col or "тег" in col or "тема" in col:
            k_idx = idx
        elif "ответ" in col or "материал" in col or "answer" in col or "информация" in col:
            a_idx = idx
        elif "отдел" in col or "dept" in col or "department" in col:
            d_idx = idx

    start_row = 1 if (k_idx != -1 or a_idx != -1) else 0
    if k_idx == -1:
        k_idx = 0
    if a_idx == -1:
        a_idx = 1 if len(raw_rows[0]) > 1 else -1

    results = []
    for row in raw_rows[start_row:]:
        if not row:
            continue
        k = str(row[k_idx] if k_idx < len(row) and row[k_idx] is not None else "").strip()
        a = str(
            row[a_idx] if a_idx != -1 and a_idx < len(row) and row[a_idx] is not None else ""
        ).strip()
        d = str(
            row[d_idx] if d_idx != -1 and d_idx < len(row) and row[d_idx] is not None else ""
        ).strip()

        if k and a:
            results.append({"keywords": k, "answer": a, "department": d})

    return results


def parse_file_or_text(
    file_bytes: bytes | None, filename: str | None, text_content: str | None = None
) -> list[list[str]]:
    """Распарсить загруженный файл (.xlsx, .csv) или текстовый ввод в матрицу строк."""
    raw_rows: list[list[str]] = []

    if file_bytes and filename:
        fn = filename.lower()
        if fn.endswith((".xlsx", ".xlsm", ".xltx")):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
            try:
                ws = wb.active
                if ws is not None:
                    raw_rows.extend(
                        [str(v or "").strip() for v in row]
                        for row in ws.iter_rows(values_only=True)
                        if any(v is not None and str(v).strip() for v in row)
                    )
            finally:
                wb.close()
            return raw_rows
        else:
            text = _decode_bytes(file_bytes)
    elif text_content:
        text = text_content
    else:
        return []

    delimiter = _detect_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows.extend([c.strip() for c in row] for row in reader if any(c.strip() for c in row))
    return raw_rows


def export_faq_xlsx(nodes: list[Any]) -> bytes:
    """Выгрузить список частых вопросов в Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Частые вопросы"

    headers = ["ID", "Отдел", "Вопрос", "Ответ"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for node in nodes:
        dept_name = node.department.name if getattr(node, "department", None) else "—"
        ws.append([node.id, dept_name, node.question, node.final_answer or ""])

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 65

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_knowledge_xlsx(items: list[Any]) -> bytes:
    """Выгрузить материалы базы знаний в Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "База знаний"

    headers = ["ID", "Отдел", "Ключевые слова", "Материал / Ответ"]
    ws.append(headers)

    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in items:
        dept_name = item.department.name if getattr(item, "department", None) else "—"
        ws.append([item.id, dept_name, item.keywords, item.answer])

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 65

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Асинхронный API: неблокирующие обёртки над синхронными ядрами.
#
# openpyxl (load_workbook / save) и csv работают только синхронно: на файлах в
# десятки тысяч строк один вызов занимает сотни миллисекунд и полностью
# останавливал event loop FastAPI (остальные запросы, longpoll и heartbeat
# бота простаивали). Каждый вызов уводится в поток пула asyncio.to_thread,
# поэтому event loop остаётся свободным.
# ---------------------------------------------------------------------------


async def generate_faq_template_xlsx_async() -> bytes:
    """Неблокирующая версия generate_faq_template_xlsx (openpyxl.save в потоке)."""
    return await asyncio.to_thread(generate_faq_template_xlsx)


async def generate_faq_template_csv_async() -> str:
    """Неблокирующая версия generate_faq_template_csv (csv.writer в потоке)."""
    return await asyncio.to_thread(generate_faq_template_csv)


async def generate_knowledge_template_xlsx_async() -> bytes:
    """Неблокирующая версия generate_knowledge_template_xlsx (openpyxl.save в потоке)."""
    return await asyncio.to_thread(generate_knowledge_template_xlsx)


async def generate_knowledge_template_csv_async() -> str:
    """Неблокирующая версия generate_knowledge_template_csv (csv.writer в потоке)."""
    return await asyncio.to_thread(generate_knowledge_template_csv)


async def parse_file_or_text_async(
    file_bytes: bytes | None, filename: str | None, text_content: str | None = None
) -> list[list[str]]:
    """Неблокирующая версия parse_file_or_text (openpyxl.load_workbook/csv.reader в потоке)."""
    return await asyncio.to_thread(parse_file_or_text, file_bytes, filename, text_content)


async def parse_faq_rows_async(raw_rows: list[list[str]]) -> list[dict[str, str]]:
    """Неблокирующая версия parse_faq_rows (разбор больших матриц вне event loop)."""
    return await asyncio.to_thread(parse_faq_rows, raw_rows)


async def parse_knowledge_rows_async(raw_rows: list[list[str]]) -> list[dict[str, str]]:
    """Неблокирующая версия parse_knowledge_rows (разбор больших матриц вне event loop)."""
    return await asyncio.to_thread(parse_knowledge_rows, raw_rows)


async def export_faq_xlsx_async(nodes: list[Any]) -> bytes:
    """Неблокирующая версия export_faq_xlsx (openpyxl.save в потоке)."""
    return await asyncio.to_thread(export_faq_xlsx, nodes)


async def export_knowledge_xlsx_async(items: list[Any]) -> bytes:
    """Неблокирующая версия export_knowledge_xlsx (openpyxl.save в потоке)."""
    return await asyncio.to_thread(export_knowledge_xlsx, items)
